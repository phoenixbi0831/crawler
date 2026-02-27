from pathlib import Path
import re


def ensure_dirs(base_dir: Path):
    doc_dir = base_dir / "doc"
    img_dir = base_dir / "img"
    doc_dir.mkdir(parents=True, exist_ok=True)
    img_dir.mkdir(parents=True, exist_ok=True)
    return doc_dir, img_dir


def to_markdown_table(rows):
    if not rows:
        return []
    md_lines = []
    header = [cell.strip() for cell in rows[0]]
    md_lines.append("| " + " | ".join(header) + " |")
    md_lines.append("| " + " | ".join("---" for _ in header) + " |")
    for row in rows[1:]:
        md_lines.append("| " + " | ".join((cell or "").strip() for cell in row) + " |")
    return md_lines


def process_docx(path: Path, doc_dir: Path, img_dir: Path):
    try:
        import docx  # type: ignore
    except ImportError:
        print(f"[WARN] 缺少 python-docx，略過：{path}")
        return

    doc = docx.Document(str(path))
    lines: list[str] = [f"# {path.stem}"]

    # 文字段落與標題
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style_name = para.style.name if para.style is not None else ""
        if style_name.startswith("Heading"):
            m = re.search(r"(\d+)", style_name)
            level = int(m.group(1)) if m else 1
            level = max(1, min(level, 3))
            lines.append(f"{'#' * level} {text}")
        else:
            lines.append(text)

    # 表格
    for table in doc.tables:
        table_rows = []
        for row in table.rows:
            table_rows.append([cell.text for cell in row.cells])
        lines.append("")
        lines.extend(to_markdown_table(table_rows))
        lines.append("")

    # 圖片
    image_counter = 1
    lines.append("")
    lines.append("## 圖片")
    rels = doc.part._rels  # type: ignore[attr-defined]
    for rel in rels.values():
        try:
            if "image" in getattr(rel._target, "content_type", ""):  # type: ignore[attr-defined]
                image = rel._target  # type: ignore[assignment]
                ext = image.content_type.split("/")[-1] or "png"  # type: ignore[attr-defined]
                img_name = f"{path.stem}_{image_counter}.{ext}"
                img_path = img_dir / img_name
                with open(img_path, "wb") as f:
                    f.write(image.blob)  # type: ignore[arg-type]
                lines.append(f"![{img_name}](../img/{img_name})")
                image_counter += 1
        except Exception as exc:  # pragma: no cover - 避免因少數關聯失敗整體中斷
            print(f"[WARN] 解析 DOCX 圖片失敗：{path} ({exc})")

    md_path = doc_dir / f"{path.stem}.md"
    md_path.write_text("\n\n".join(lines), encoding="utf-8")
    print(f"[OK] DOCX -> {md_path}")


def process_pptx(path: Path, doc_dir: Path, img_dir: Path):
    try:
        from pptx import Presentation  # type: ignore
    except ImportError:
        print(f"[WARN] 缺少 python-pptx，略過：{path}")
        return

    prs = Presentation(str(path))
    lines: list[str] = [f"# {path.stem}"]
    image_counter = 1

    for slide_index, slide in enumerate(prs.slides, start=1):
        lines.append(f"## 投影片 {slide_index}")
        for shape in slide.shapes:
            # 文字
            if getattr(shape, "has_text_frame", False):
                text_chunks = []
                for para in shape.text_frame.paragraphs:  # type: ignore[attr-defined]
                    if para.text.strip():
                        text_chunks.append(para.text.strip())
                if text_chunks:
                    # 第一行當作小標，其餘視作內文
                    first, *rest = text_chunks
                    lines.append(f"### {first}")
                    for t in rest:
                        lines.append(t)

            # 表格
            if getattr(shape, "has_table", False):
                table = shape.table  # type: ignore[attr-defined]
                table_rows = []
                for row in table.rows:
                    table_rows.append([cell.text for cell in row.cells])
                lines.append("")
                lines.extend(to_markdown_table(table_rows))
                lines.append("")

            # 圖片
            image = getattr(shape, "image", None)
            if image is not None:
                ext = getattr(image, "ext", "png") or "png"
                img_name = f"{path.stem}_slide{slide_index}_{image_counter}.{ext}"
                img_path = img_dir / img_name
                with open(img_path, "wb") as f:
                    f.write(image.blob)  # type: ignore[arg-type]
                lines.append(f"![{img_name}](../img/{img_name})")
                image_counter += 1

    md_path = doc_dir / f"{path.stem}.md"
    md_path.write_text("\n\n".join(lines), encoding="utf-8")
    print(f"[OK] PPTX -> {md_path}")


def process_pdf(path: Path, doc_dir: Path, img_dir: Path):  # img_dir 保留給日後需要抽圖片時用
    try:
        import pdfplumber  # type: ignore
    except ImportError:
        print(f"[WARN] 缺少 pdfplumber，略過：{path}")
        return

    lines: list[str] = [f"# {path.stem}"]
    with pdfplumber.open(str(path)) as pdf:
        for page_index, page in enumerate(pdf.pages, start=1):
            lines.append(f"## 第 {page_index} 頁")
            text = page.extract_text() or ""
            text = text.strip()
            if text:
                lines.append(text)

            tables = page.extract_tables() or []
            for table in tables:
                lines.append("")
                lines.extend(to_markdown_table(table))
                lines.append("")

    md_path = doc_dir / f"{path.stem}.md"
    md_path.write_text("\n\n".join(lines), encoding="utf-8")
    print(f"[OK] PDF -> {md_path}")


def main():
    base_dir = Path(__file__).resolve().parent
    resource_dir = base_dir / "resource"
    if not resource_dir.exists():
        print(f"[ERROR] 找不到 resource 資料夾：{resource_dir}")
        return

    doc_dir, img_dir = ensure_dirs(base_dir)

    for path in resource_dir.iterdir():
        if not path.is_file():
            continue
        ext = path.suffix.lower()
        if ext == ".docx":
            process_docx(path, doc_dir, img_dir)
        elif ext == ".pptx":
            process_pptx(path, doc_dir, img_dir)
        elif ext == ".pdf":
            process_pdf(path, doc_dir, img_dir)
        else:
            # 其他格式先略過，有需要再補
            print(f"[SKIP] 不支援的副檔名：{path.name}")


if __name__ == "__main__":
    main()

